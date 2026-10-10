# -*- coding: utf-8 -*-
"""
把已有的 xlsx 底稿抽成网页项目统一的机器可读 schema。
只读原文件, 不改动任何原有资产。
输出: data/519191.json, data/002910.json
"""
import json
import os
import re
import openpyxl

LIB = r"C:\Users\75719\Desktop\基金复盘_小组作业\02_现有成果\基金复盘库"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data")
os.makedirs(OUT, exist_ok=True)


def num(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("%", "").replace(",", "")
    if s in ("", "---", "--", "—", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def quarter_period(label):
    """'2020Q4(2020-12-31)' -> ('2020Q4', '2020-12-31')"""
    m = re.match(r"\s*(\d{4}Q[1-4])\s*[（(]?\s*([\d\-]*)\s*[)）]?", str(label))
    if not m:
        return None, None
    return m.group(1), (m.group(2) or "")


# ---------------------------------------------------------------- 519191
def load_519191():
    p = os.path.join(LIB, "万家新利灵活配置混合_519191_前十大持仓与历史数据.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)

    # --- 基金概览
    fund = {"code": "519191"}
    ws = wb["基金概览"]
    for r in ws.iter_rows(values_only=True):
        k, v = (r + ("", ""))[:2]
        if k:
            fund[str(k).strip()] = str(v).strip() if v is not None else ""

    # --- 季度汇总(集中度/主题/规模/仓位/净值)
    summary = {}
    ws = wb["季度汇总"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    for r in rows:
        per = str(r[0]).strip() if r[0] else ""
        if not re.match(r"\d{4}Q[1-4]", per):
            continue
        summary[per] = {
            "top10Sum": num(r[1]),
            "top5Sum": num(r[2]),
            "theme": (str(r[3]).strip() if r[3] else ""),
            "themeWeight": num(r[4]),
            "shares": num(r[5]),
            "aum": num(r[6]),
            "aumChange": (str(r[7]).strip() if r[7] else ""),
            "stockPos": num(r[8]),
            "nav": num(r[9]),
        }

    # --- 持仓变动(留存/Jaccard/新进/退出)
    churn = {}
    ws = wb["持仓变动"]
    for r in ws.iter_rows(min_row=3, values_only=True):
        per = str(r[0]).strip() if r[0] else ""
        if not re.match(r"\d{4}Q[1-4]", per):
            continue
        churn[per] = {
            "retained": num(r[2]),
            "jaccard": num(r[3]),
            "added": [s for s in re.split(r"[、,]", str(r[4] or "")) if s.strip() and s.strip() != "—"],
            "dropped": [s for s in re.split(r"[、,]", str(r[5] or "")) if s.strip() and s.strip() != "—"],
        }

    # --- 历史前十长表
    hist = {}
    ws = wb["历史前十"]
    for r in ws.iter_rows(min_row=3, values_only=True):
        per, end = quarter_period(r[0])
        if not per:
            continue
        hist.setdefault(per, {"endDate": end, "top10": []})
        hist[per]["top10"].append({
            "rank": int(num(r[1]) or 0),
            "code": (str(r[2]).strip().zfill(6) if r[2] else ""),
            "name": (str(r[3]).strip() if r[3] else ""),
            "weight": num(r[4]),
            "sharesWan": num(r[5]),
            "valueWan": num(r[6]),
        })

    # --- 资产与规模序列
    assets = {}
    ws = wb["资产与规模"]
    for r in ws.iter_rows(min_row=3, values_only=True):
        per = str(r[0]).strip() if r[0] else ""
        if not re.match(r"\d{4}Q[1-4]", per):
            continue
        assets[per] = {"stockPos": num(r[1]), "cashPos": num(r[3]), "aum": num(r[4])}

    periods = sorted(hist.keys())
    review = [p for p in periods if re.match(r"20(2[1-9]|3\d)", p[2:]) or True]
    # 复盘区间 = 2020Q4 起(2020Q3 为前任末, 标为参考)
    review = [p for p in periods if p >= "2020Q4"]

    quarters = []
    for p in review:
        h = hist[p]
        s = summary.get(p, {})
        c = churn.get(p, {})
        a = assets.get(p, {})
        quarters.append({
            "period": p,
            "endDate": h["endDate"],
            "top10": sorted(h["top10"], key=lambda x: x["rank"]),
            "top10Sum": s.get("top10Sum"),
            "top5Sum": s.get("top5Sum"),
            "theme": s.get("theme") or "",
            "themeWeight": s.get("themeWeight"),
            "aum": s.get("aum") if s.get("aum") is not None else a.get("aum"),
            "shares": s.get("shares"),
            "nav": s.get("nav"),
            "stockPos": s.get("stockPos") if s.get("stockPos") is not None else a.get("stockPos"),
            "aumChange": s.get("aumChange") or "",
            "retained": c.get("retained"),
            "jaccard": c.get("jaccard"),
            "added": c.get("added", []),
            "dropped": c.get("dropped", []),
        })

    prev = hist.get("2020Q3")
    return {
        "schema": "fundclinic/1",
        "fund": {
            "code": "519191",
            "name": fund.get("基金简称", "万家新利灵活配置混合"),
            "fullName": fund.get("基金全称", ""),
            "company": fund.get("基金公司", ""),
            "type": fund.get("基金类型", ""),
            "manager": "黄海",
            "managerSince": "2020-09-24 接手",
            "range": "2020Q4–2026Q2",
            "quarterCount": len(quarters),
        },
        "perf": {
            "cumReturn": 76.77,
            "exBestQuarter": 35.47,
            "bestQuarter": "2022Q1",
            "maxDrawdown": -32.45,
            "aumPeak": 16.72,
            "aumEnd": 2.92,
            "note": "区间复利收益为官方季报口径; 剔除 2022Q1 单季爆发的口径见复盘正文",
        },
        "prevTenure": {
            "period": "2020Q3",
            "endDate": prev["endDate"] if prev else "",
            "top10": sorted(prev["top10"], key=lambda x: x["rank"]) if prev else [],
        },
        "quarters": quarters,
        "sources": [
            "官方定期报告 PDF: 巨潮资讯 static.cninfo.com.cn / 证监会 eid.csrc.gov.cn / 东财镜像 pdf.dfcfw.com(23/23 期已取得)",
            "天天基金 F10 历史前十大持仓 fundf10.eastmoney.com/ccmx_519191.html",
            "天天基金 规模变动 fundf10.eastmoney.com/gmbd_519191.html",
            "天天基金 净值序列 fund.eastmoney.com/pingzhongdata/519191.js",
        ],
    }


# ---------------------------------------------------------------- 002910
def load_002910():
    p = os.path.join(LIB, "易方达供给改革混合_002910_持仓全历史_2022-2026.xlsx")
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb["持仓变动矩阵"]
    header = None
    matrix = {}   # period -> list
    for r in ws.iter_rows(min_row=3, values_only=True):
        if header is None:
            header = [str(x).strip() if x else "" for x in r]
            periods = [h for h in header[2:] if re.match(r"\d{4}Q[1-4]", h)]
            continue
        code = str(r[0]).strip().zfill(6) if r[0] else ""
        name = str(r[1]).strip() if r[1] else ""
        if not code or not name:
            continue
        for i, per in enumerate(periods):
            w = num(r[2 + i]) if (2 + i) < len(r) else None
            if w is None:
                continue
            matrix.setdefault(per, []).append({"code": code, "name": name, "weight": w})

    quarters = []
    for per in periods:
        lst = sorted(matrix.get(per, []), key=lambda x: -x["weight"])[:10]
        top10 = [{"rank": i + 1, "code": x["code"], "name": x["name"], "weight": x["weight"]} for i, x in enumerate(lst)]
        quarters.append({
            "period": per,
            "endDate": "",
            "top10": top10,
            "top10Sum": round(sum(x["weight"] for x in top10), 2) if top10 else None,
            "top5Sum": round(sum(x["weight"] for x in top10[:5]), 2) if top10 else None,
            "theme": "",
            "themeWeight": None,
            "aum": None, "shares": None, "nav": None, "stockPos": None, "aumChange": "",
            "retained": None, "jaccard": None, "added": [], "dropped": [],
        })

    return {
        "schema": "fundclinic/1",
        "fund": {
            "code": "002910",
            "name": "易方达供给改革混合",
            "fullName": "易方达供给改革灵活配置混合型证券投资基金",
            "company": "易方达基金管理有限公司",
            "type": "混合型-灵活配置",
            "manager": "杨宗昌",
            "managerSince": "2019-01 起管理",
            "range": "2022Q1–2026Q2",
            "quarterCount": len(quarters),
        },
        "perf": None,
        "prevTenure": None,
        "quarters": quarters,
        "sources": [
            "天天基金 F10 历史前十大持仓 fundf10.eastmoney.com/ccmx_002910.html",
            "官方定期报告 PDF(基金公司官网 / 巨潮资讯)",
            "注: 本样例仅含持仓结构数据, 未附净值与收益序列",
        ],
    }


if __name__ == "__main__":
    bundle = {}
    for fn, loader in (("519191.json", load_519191), ("002910.json", load_002910)):
        d = loader()
        fp = os.path.join(OUT, fn)
        with open(fp, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, separators=(",", ":"))
        bundle[d["fund"]["code"]] = d
        q = d["quarters"]
        print(fn, "quarters=", len(q), "bytes=", os.path.getsize(fp),
              "range=", q[0]["period"], "->", q[-1]["period"])
        print("   latest top10:", ", ".join(f'{x["name"]}{x["weight"]}' for x in q[-1]["top10"][:4]))

    # 网页项目直接内联(要能在 file:// 双击打开, 不能用 fetch)
    fp = os.path.join(OUT, "funds.data.js")
    with open(fp, "w", encoding="utf-8") as f:
        f.write("/* 由 tools/extract_data.py 生成, 请勿手改。来源: 公开季报 / 天天基金 F10 */\n")
        f.write("window.FUNDCLINIC_DATA = ")
        json.dump(bundle, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print("funds.data.js bytes=", os.path.getsize(fp))
